import subprocess
import sys
subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=False)

import os
import uuid
import asyncio
import sys

if sys.platform == 'win32':
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

from typing import Dict, Any, List
from fastapi import FastAPI, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import traceback
import io

# openpyxl for premium Excel formatting
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scrapers.simulated import scrape_simulated
from scrapers.real_playwright import scrape_real_playwright

app = FastAPI(title="MapLead Outscraper Clone Backend")

# Setup CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# In-memory Job Database
jobs: Dict[str, Dict[str, Any]] = {}

class ScrapeRequest(BaseModel):
    category: str
    area: str
    limit: int = 10
    engine: str = "simulated"

@app.get("/")
def read_root():
    # Serves the gorgeous frontend page
    static_index = os.path.join(os.path.dirname(__file__), "static", "index.html")
    if os.path.exists(static_index):
        return FileResponse(static_index)
    return {"message": "MapLead Python Backend Active. Static index.html not found."}

@app.post("/api/scrape")
async def start_scrape(request: ScrapeRequest, background_tasks: BackgroundTasks):
    if not request.category or not request.area:
        raise HTTPException(status_code=400, detail="Category and Area are required fields.")
        
    # Respect the client's requested limit exactly (use as provided).
    # If client omits limit, pydantic default (10) will apply.
    limit = int(request.limit)
    engine = "real"

    job_id = f"job_{int(uuid.uuid4().time_low)}_{engine}"

    jobs[job_id] = {
        "id": job_id,
        "category": request.category,
        "area": request.area,
        "limit": limit,
        "engine": engine,
        "status": "running",
        "progress": 0,
        "logs": [f"[INIT] Python job created: Scrape '{request.category}' in '{request.area}' (Limit: {limit}, Engine: {engine.upper()})"],
        "results": [],
    }

    # Dispatch to asyncio background tasks
    background_tasks.add_task(run_scraper_job, job_id, request.category, request.area, limit, engine)

    return {"jobId": job_id, "message": "Scraping task dispatched successfully."}

@app.get("/api/scrape/status/{job_id}")
async def get_status(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found.")
    
    job = jobs[job_id]
    return {
        "id": job["id"],
        "status": job["status"],
        "progress": job["progress"],
        "logs": job["logs"],
        "results": job["results"] if job["status"] == "completed" else [],
        "error": job.get("error")
    }

@app.get("/api/download/{job_id}")
async def download_excel(job_id: str):
    job = jobs.get(job_id)
    if not job:
        # Fallback: allow download when engine suffix is missing or job key was altered slightly
        fallback_key = next((key for key in jobs if key.startswith(job_id)), None)
        if fallback_key:
            job = jobs[fallback_key]

    if not job:
        raise HTTPException(status_code=404, detail="Job not found.")
        
    if job["status"] != "completed":
        raise HTTPException(status_code=400, detail="Scrape report is not compiled yet.")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Scraped Leads"

        outscraper_headers = [
            'query', 'name', 'name_for_emails', 'subtypes', 'category', 'type', 'phone', 
            'phone.phones_enricher.carrier_name', 'phone.phones_enricher.carrier_type', 
            'website', 'address', 'street', 'city', 'county', 'state', 'state_code', 
            'postal_code', 'country', 'country_code', 'domain', 'company_name', 'company_phone', 
            'company_phone.phones_enricher.carrier_name', 'company_phone.phones_enricher.carrier_type', 
            'company_phones', 'company_linkedin', 'company_facebook', 'company_instagram', 'company_x', 
            'company_youtube', 'full_name', 'first_name', 'last_name', 'title', 'email', 
            'email.emails_validator.status', 'email.emails_validator.status_details', 'contact_phone', 
            'contact_phones', 'contact_linkedin', 'contact_facebook', 'contact_instagram', 'contact_x', 
            'website_title', 'website_description', 'website_generator', 'website_has_gtm', 
            'website_has_fb_pixel', 'source', 'h3', 'time_zone', 
            'plus_code', 'area_service', 'rating', 'reviews', 'reviews_link', 'reviews_tags', 
            'reviews_per_score', 'reviews_per_score_1', 'reviews_per_score_2', 'reviews_per_score_3', 
            'reviews_per_score_4', 'reviews_per_score_5', 'photos_count', 'photo', 'street_view', 
            'logo', 'located_in', 'located_google_id', 'business_status', 'working_hours', 
            'working_hours_csv_compatible', 'other_hours', 'popular_times', 'typical_time_spent', 
            'range', 'prices', 'reservation_links', 'booking_appointment_link', 'menu_link', 
            'order_links', 'about', 'description', 'posts', 'verified', 'owner_id', 'owner_title', 
            'owner_link', 'location_link', 'location_reviews_link', 'place_id', 'google_id', 'cid', 
            'kgmid', 'reviews_id', 'company_insights.employees', 'company_insights.revenue', 
            'company_insights.founded_year', 'company_insights.industry', 'company_insights.is_public', 
            'company_insights.name', 'company_insights.country', 'company_insights.state', 
            'company_insights.city', 'company_insights.zip', 'company_insights.address', 'chain_info.chain'
        ]

        # Determine which headers actually contain data and exclude empty columns
        def _map_value(h, item):
            if h == 'query':
                return f"{job['category']} in {job['area']}"
            if h == 'name':
                return item.get('name', '')
            if h == 'name_for_emails':
                return item.get('name', '')
            if h == 'category':
                return item.get('category', '')
            if h == 'type':
                return item.get('category', '')
            if h == 'phone':
                val = item.get('phone', '')
                return '' if val in [None, 'N/A', 'Phone not listed'] else val
            if h == 'website':
                return item.get('website', '') or ''
            if h == 'address':
                val = item.get('address', '')
                return '' if val in [None, 'Extraction error', 'Address not listed'] else val
            if h == 'email':
                return item.get('email', '')
            if h == 'company_facebook':
                return item.get('socials', {}).get('facebook', '')
            if h == 'company_instagram':
                return item.get('socials', {}).get('instagram', '')
            if h == 'latitude':
                return item.get('latitude', '')
            if h == 'longitude':
                return item.get('longitude', '')
            if h == 'rating':
                return item.get('rating', '') or ''
            if h == 'reviews':
                return item.get('reviewsCount', '') or ''
            # Default: try to fetch key directly
            return item.get(h, '') if isinstance(item, dict) else ''

        selected_headers = []
        for h in outscraper_headers:
            has_value = False
            for it in job.get('results', []):
                try:
                    v = _map_value(h, it)
                except Exception:
                    v = ''
                if v not in (None, '', 'N/A', 'No website listed', 'Phone not listed', 'Address not listed'):
                    has_value = True
                    break
            if has_value or h in ('query', 'name'):
                selected_headers.append(h)

        if not selected_headers:
            selected_headers = ['query', 'name']

        # Write Headers
        font_header = Font(bold=True)
        for col_idx, h in enumerate(selected_headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = font_header

        # Write Data
        for index, item in enumerate(job.get('results', [])):
            row_idx = index + 2
            for col_idx, h in enumerate(selected_headers, 1):
                val = _map_value(h, item)

                # Ensure proper URL formatting for links
                if h in ['website', 'company_facebook', 'company_instagram', 'owner_link', 'location_link']:
                    url_val = str(val).strip()
                    if url_val and url_val.lower() not in ['n/a', 'no website listed']:
                        if not url_val.startswith('http'):
                            url_val = 'https://' + url_val
                        val = url_val
                    else:
                        val = ""

                # Set clickable hyperlink for URLs natively
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if isinstance(val, str) and val.startswith("http"):
                    try:
                        cell.hyperlink = val
                        cell.font = Font(color="0563C1", underline="single")
                    except ValueError:
                        pass

        # Stream compiled spreadsheet
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        safe_category = job["category"].replace(" ", "_")
        safe_area = job["area"].replace(" ", "_")
        filename = f"Outscraper-Tasks-{safe_category}_{safe_area}.xlsx"

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Excel generation failed: {str(e)}")

async def run_scraper_job(job_id: str, category: str, area: str, limit: int, engine: str):
    async def progress_update(progress: int, log_line: str):
        if job_id in jobs:
            jobs[job_id]["progress"] = progress
            jobs[job_id]["logs"].append(log_line)
            print(f"[Task {job_id} - {progress}%] {log_line}")

    try:
        if engine == "real":
            scraped_data = await scrape_real_playwright(category, area, limit, progress_update)
        else:
            scraped_data = await scrape_simulated(category, area, limit, progress_update)

        if job_id in jobs:
            jobs[job_id]["status"] = "completed"
            jobs[job_id]["progress"] = 100
            jobs[job_id]["results"] = scraped_data

    except Exception as error:
        tb = traceback.format_exc()
        print(f"Scraper Job failed: {repr(error)}")
        print(tb)
        if job_id in jobs:
            jobs[job_id]["status"] = "failed"
            jobs[job_id]["error"] = repr(error)
            jobs[job_id]["logs"].append(f"[CRITICAL ERROR] Pipeline crashed: {repr(error)}")
            jobs[job_id]["logs"].append(f"[DEBUG] Traceback:\n{tb}")

# Serve Static files using a simple wrapper
# FastAPI can serve directory structures natively
from fastapi.staticfiles import StaticFiles
static_path = os.path.join(os.path.dirname(__file__), "static")
if os.path.exists(static_path):
    app.mount("/static", StaticFiles(directory=static_path), name="static")
