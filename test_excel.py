import sys
import os
sys.path.insert(0, os.path.join(os.getcwd(), 'outscraper-python'))
from app import jobs
from openpyxl import Workbook
from openpyxl.styles import Font

# Mock job
job = {
    "status": "completed",
    "category": "Travel agent",
    "area": "Satellite, Ahmedabad",
    "results": [
        {
            "id": 1,
            "name": "Bhagwati Aviations",
            "category": "Travel agent",
            "phone": "098250 30555",
            "website": "http://www.bhagwatiholidays.com/",
            "email": "info@bhagwatiholidays.com",
            "socials": {"facebook": "", "instagram": ""},
            "latitude": 0, "longitude": 0, "rating": 4.8, "reviewsCount": 90, "address": "Test"
        }
    ]
}

outscraper_headers = ['query', 'name', 'website', 'email', 'phone']

wb = Workbook()
ws = wb.active
for col_idx, h in enumerate(outscraper_headers, 1):
    ws.cell(row=1, column=col_idx, value=h)

for index, item in enumerate(job["results"]):
    row_idx = index + 2
    for col_idx, h in enumerate(outscraper_headers, 1):
        val = item.get(h, "")
        
        if h in ['website']:
            url_val = str(val).strip()
            if url_val and not url_val.startswith('http'):
                url_val = 'https://' + url_val
            val = url_val

        # Set clickable hyperlink for URLs using Excel Formula
        if isinstance(val, str) and val.startswith("http"):
            safe_url = val.replace('"', '""')
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = f'=HYPERLINK("{safe_url}", "{safe_url}")'
            cell.font = Font(color="0563C1", underline="single")
        else:
            cell = ws.cell(row=row_idx, column=col_idx, value=val)

wb.save('test_excel.xlsx')
print("Saved test_excel.xlsx")
